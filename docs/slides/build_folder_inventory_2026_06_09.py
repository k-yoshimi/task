"""TASK リポジトリ フォルダ棚卸しスライド (日本語, 2026-06-09)。

実行方法::

    python3 docs/slides/build_folder_inventory_2026_06_09.py

出力: docs/slides/2026-06-09-task-folder-inventory.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from build_slides import (
    SLIDE_W,
    SLIDE_H,
    COLOR_TITLE,
    COLOR_ACCENT,
    COLOR_TEXT,
    COLOR_DIM,
    add_title_bar,
    add_text_block,
    add_table,
    set_jp_font,
)

DECK_TAG = "TASK リポジトリ フォルダ棚卸し 2026-06-09"


def add_deck_footer(slide, page_no, total_pages):
    box = slide.shapes.add_textbox(
        Inches(0.3), Inches(7.05), Inches(12.7), Inches(0.35)
    )
    tf = box.text_frame
    tf.margin_top = Inches(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = f"{DECK_TAG} — {page_no} / {total_pages}"
    set_jp_font(r, size=10, color=COLOR_DIM)


def add_bullets_v2(slide, bullets, top, left=Inches(0.6),
                   width=Inches(12.1), height=Inches(5.6),
                   size=18, vertical_anchor=None):
    """空文字列を「バレット記号無しの小さなスペーサ段落」として扱うバレット描画。"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if vertical_anchor is not None:
        tf.vertical_anchor = vertical_anchor
    first = True
    for item in bullets:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
            if text.startswith("  - "):
                level = 1
                text = text[4:]
            elif text.startswith("- "):
                text = text[2:]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.alignment = PP_ALIGN.LEFT
        if not text.strip():
            r = p.add_run()
            r.text = " "
            set_jp_font(r, size=8)
            continue
        prefix = "  " * level + ("• " if level == 0 else "– ")
        r = p.add_run()
        r.text = prefix + text
        set_jp_font(r, size=size - 2 * level)


# ---- スライド本体 ----------------------------------------------------------


def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(1.9), SLIDE_W, Inches(2.7)
    )
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_TITLE
    add_text_block(
        s,
        ["TASK リポジトリ",
         "フォルダ棚卸し & 整理方針"],
        top=Inches(2.2), left=Inches(0.8), width=Inches(11.7),
        height=Inches(2.0), size=38,
        color=RGBColor(0xFF, 0xFF, 0xFF), bold=True,
    )
    add_text_block(
        s,
        ["— まず棚卸し: 約60個のトップレベルフォルダを分類。物理移動はしない —"],
        top=Inches(4.8), left=Inches(0.8), width=Inches(11.7),
        height=Inches(0.6), size=16, color=COLOR_ACCENT,
    )
    line = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(5.6), Inches(2.5), Inches(0.05),
    )
    line.line.fill.background()
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_ACCENT
    add_text_block(
        s,
        ["吉見一慶 (東京大学)",
         "2026-06-09"],
        top=Inches(5.8), left=Inches(0.8), width=Inches(11.7),
        height=Inches(1.0), size=14,
        color=RGBColor(0x55, 0x55, 0x55),
    )
    return s


def slide_scope(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(s, "1. 目的とスコープ")
    add_bullets_v2(s, [
        "目的: リポジトリ約60個のトップレベルフォルダを棚卸しし、現状を分類する。",
        "本資料は何も移動・削除しない。棚卸しのみ。",
        "",
        "背景: ツリーには4つの「世界」が混在している:",
        "  - 進行中の Python 化 / リファクタリング・プロジェクトの資産,",
        "  - 継承した上流 Fortran モジュール,",
        "  - 個人 / 実験用のミラーフォルダ,",
        "  - 複数系統に分かれて増殖したドキュメント。",
        "",
        "重要な制約: 上流 Fortran モジュール本体 (eq / tr / fp / wr / wm / wf … )",
        "  - の物理的な移動・削除は ats-fukuyama さまのサインオフが必須。",
        "  - C-ABI / ラッパー / レジストリ / テスト / docs はこの制約の対象外。",
        "",
        "今回のスコープ外: あらゆる git mv / 削除、グループ C の個別判定 (TBC のまま)。",
    ], top=Inches(1.1), size=16, vertical_anchor=MSO_ANCHOR.MIDDLE, height=Inches(5.8))
    return s


def slide_overview(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(s, "2. 全体像 — 約60フォルダ → 4分類")
    add_table(s,
        ["分類", "意味", "移動方針"],
        [
            ["A — プロジェクト所有・現役",
             "Phase L モジュール + Python ラッパー + MCP + プロジェクト docs",
             "自由に編集可 (自分たちの資産)"],
            ["B — 上流現役",
             "ルート README 記載のモジュール。継承物",
             "ats-fukuyama サインオフ必須"],
            ["C — バリアント / ミラー",
             "個人名付き・派生フォルダ。アーカイブ候補",
             "TBC — ユーザーが分類"],
            ["D — ドキュメント整理",
             "重複した doc 根、ビルド生成物、薄いフォルダ、散乱物",
             "docs/ 内で整理 (自分たちの資産)"],
        ],
        top=Inches(1.4),
        col_widths=[Inches(3.4), Inches(5.7), Inches(3.0)],
        row_height=Inches(0.85),
    )
    add_text_block(s, [
        "以降のスライドで使う凡例。「最終更新」= 最終 git コミット日。",
        "多くは 2026-04-17/18 (一括インポート) で、それ以降の日付が実際に編集したフォルダの印。",
    ], top=Inches(5.6), size=13, color=COLOR_ACCENT)
    return s


# ---- メイン ----------------------------------------------------------------


BUILDERS = [
    slide_title,
    slide_scope,
    slide_overview,
]


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    total = len(BUILDERS)
    for i, builder in enumerate(BUILDERS, start=1):
        s = builder(prs)
        if i > 1:  # 表紙にはフッタを付けない
            add_deck_footer(s, i, total)

    out = Path(__file__).parent / "2026-06-09-task-folder-inventory.pptx"
    prs.save(out)
    print(f"Wrote {out} ({total} slides)")
    return out


if __name__ == "__main__":
    build()
