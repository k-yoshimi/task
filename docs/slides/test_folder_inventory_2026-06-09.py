"""Smoke test for the folder-inventory deck build script."""
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))  # so `import build_slides` inside the script resolves

import build_folder_inventory_2026_06_09 as deck  # noqa: E402

# Bumped as slides are added, task by task.
EXPECTED_SLIDES = 9


def test_build_produces_deck():
    out = deck.build()
    assert out.exists(), f"deck not written: {out}"
    prs = Presentation(str(out))
    assert len(prs.slides) == EXPECTED_SLIDES


def test_no_placeholder_text():
    # Group C dispositions are intentionally "TBC"/"要確認"; no other placeholders.
    src = (HERE / "build_folder_inventory_2026_06_09.py").read_text()
    for bad in ("TODO", "FIXME", "XXX_PLACEHOLDER"):
        assert bad not in src, f"placeholder {bad!r} left in build script"
